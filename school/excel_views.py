from io import BytesIO
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .models import ClassLevel, ClassSubject, SchoolProfile
from .result_views import build_class_results, get_current_term


@login_required(login_url="/teacher-login/")
def class_master_result_excel(request, class_id):
    class_level = get_object_or_404(ClassLevel, id=class_id)
    term = get_current_term()

    if term is None:
        return HttpResponse(
            "No academic term has been configured.",
            status=400,
        )

    result = build_class_results(class_level, term)
    students = result.get("students", [])

    # Use ClassSubject so every subject assigned to the class
    # appears, even when no score has been entered.
    class_subjects = (
        ClassSubject.objects
        .filter(
            class_level=class_level,
            active=True,
        )
        .select_related("subject")
        .order_by("subject__name")
    )

    subjects = [
        {
            "subject_id": cs.subject_id,
            "code": cs.subject.code,
            "name": cs.subject.name,
        }
        for cs in class_subjects
    ]

    school = SchoolProfile.objects.first()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Master Result"

    # ---------------------------------------------------------
    # STYLES
    # ---------------------------------------------------------

    green_fill = PatternFill(
        fill_type="solid",
        fgColor="0B6B3A",
    )

    light_green_fill = PatternFill(
        fill_type="solid",
        fgColor="EAF3ED",
    )

    white_font = Font(
        color="FFFFFF",
        bold=True,
    )

    school_font = Font(
        bold=True,
        size=16,
    )

    title_font = Font(
        bold=True,
        size=14,
    )

    bold_font = Font(
        bold=True,
    )

    thin = Side(
        style="thin",
        color="808080",
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )

    center = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )

    left = Alignment(
        horizontal="left",
        vertical="center",
        wrap_text=True,
    )

    # ---------------------------------------------------------
    # SCHOOL HEADER
    # ---------------------------------------------------------

    last_column = 3 + len(subjects) + 4

    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=last_column,
    )

    worksheet.cell(
        1,
        1,
        school.name if school else "SCHOOL NAME",
    )

    worksheet.cell(1, 1).font = school_font
    worksheet.cell(1, 1).alignment = center

    row = 2

    if school and school.address:
        worksheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=last_column,
        )
        worksheet.cell(row, 1, school.address)
        worksheet.cell(row, 1).alignment = center
        row += 1

    contact = ""

    if school:
        contact_parts = []

        if school.phone:
            contact_parts.append(
                f"Tel: {school.phone}"
            )

        if school.email:
            contact_parts.append(
                f"Email: {school.email}"
            )

        contact = " | ".join(contact_parts)

    if contact:
        worksheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=last_column,
        )
        worksheet.cell(row, 1, contact)
        worksheet.cell(row, 1).alignment = center
        row += 1

    if school and school.motto:
        worksheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=last_column,
        )
        worksheet.cell(
            row,
            1,
            school.motto,
        )
        worksheet.cell(row, 1).alignment = center
        worksheet.cell(row, 1).font = Font(
            italic=True,
        )
        row += 2

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=last_column,
    )

    worksheet.cell(
        row,
        1,
        "MASTER RESULT",
    )

    worksheet.cell(row, 1).font = title_font
    worksheet.cell(row, 1).alignment = center

    row += 1

    worksheet.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=last_column,
    )

    worksheet.cell(
        row,
        1,
        f"CLASS: {class_level.name}    |    "
        f"SESSION: {term.session.name}    |    "
        f"TERM: {term.name}",
    )

    worksheet.cell(row, 1).font = bold_font
    worksheet.cell(row, 1).alignment = center

    row += 2

    # ---------------------------------------------------------
    # TABLE HEADER
    # ---------------------------------------------------------

    header_row = row

    headers = [
        "Position",
        "Admission No.",
        "Student Name",
    ]

    for subject in subjects:
        headers.append(
            f"{subject['name']}\n"
            "CA1 | CA2 | CA3 | Exam | Total"
        )

    headers.extend(
        [
            "Total Marks",
            "Average",
            "Grade",
            "Remark",
        ]
    )

    for column, value in enumerate(
        headers,
        start=1,
    ):
        cell = worksheet.cell(
            header_row,
            column,
            value,
        )

        cell.fill = green_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border

    row += 1

    # ---------------------------------------------------------
    # STUDENT DATA
    # ---------------------------------------------------------

    for student in students:

        row_data = [
            student.get("position", ""),
            student["student"].admission_number,
            student["student"].full_name,
        ]

        subject_rows = student.get(
            "subjects",
            [],
        )

        for subject in subjects:

            matching = next(
                (
                    subject_row
                    for subject_row in subject_rows
                    if isinstance(
                        subject_row.get("subject"),
                        dict,
                    )
                    and subject_row["subject"].get(
                        "subject_id"
                    ) == subject["subject_id"]
                ),
                None,
            )

            if matching is None:
                row_data.append(
                    "CA1: -\n"
                    "CA2: -\n"
                    "CA3: -\n"
                    "Exam: -\n"
                    "Total: -"
                )
                continue

            def fmt(value):
                if value is None:
                    return "-"
                return f"{Decimal(str(value)):.2f}"

            row_data.append(
                f"CA1: {fmt(matching.get('ca1'))}\n"
                f"CA2: {fmt(matching.get('ca2'))}\n"
                f"CA3: {fmt(matching.get('ca3'))}\n"
                f"Exam: {fmt(matching.get('exam'))}\n"
                f"Total: {fmt(matching.get('total'))}"
            )

        total_marks = student.get(
            "total_marks",
            "",
        )

        average = student.get(
            "average",
            "",
        )

        if total_marks not in ("", None):
            total_marks = float(
                Decimal(str(total_marks))
            )

        if average not in ("", None):
            average = float(
                Decimal(str(average))
            )

        row_data.extend(
            [
                total_marks,
                average,
                student.get(
                    "overall_grade",
                    "",
                ),
                student.get(
                    "overall_remark",
                    "",
                ),
            ]
        )

        for column, value in enumerate(
            row_data,
            start=1,
        ):
            cell = worksheet.cell(
                row,
                column,
                value,
            )

            cell.border = border
            cell.alignment = (
                left
                if column == 3
                else center
            )

        row += 1

    # ---------------------------------------------------------
    # EMPTY CLASS
    # ---------------------------------------------------------

    if not students:
        worksheet.merge_cells(
            start_row=row,
            start_column=1,
            end_row=row,
            end_column=last_column,
        )

        worksheet.cell(
            row,
            1,
            "No students found for this class and term.",
        )

        worksheet.cell(
            row,
            1,
        ).alignment = center

    # ---------------------------------------------------------
    # COLUMN WIDTHS
    # ---------------------------------------------------------

    worksheet.column_dimensions["A"].width = 11
    worksheet.column_dimensions["B"].width = 17
    worksheet.column_dimensions["C"].width = 28

    for column in range(
        4,
        4 + len(subjects),
    ):
        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = 20

    final_start = 4 + len(subjects)

    worksheet.column_dimensions[
        get_column_letter(final_start)
    ].width = 14

    worksheet.column_dimensions[
        get_column_letter(final_start + 1)
    ].width = 12

    worksheet.column_dimensions[
        get_column_letter(final_start + 2)
    ].width = 10

    worksheet.column_dimensions[
        get_column_letter(final_start + 3)
    ].width = 25

    # ---------------------------------------------------------
    # FREEZE PANES / FILTER / PRINT SETTINGS
    # ---------------------------------------------------------

    worksheet.freeze_panes = (
        f"D{header_row + 1}"
    )

    worksheet.auto_filter.ref = (
        f"A{header_row}:"
        f"{get_column_letter(last_column)}"
        f"{header_row + len(students)}"
    )

    worksheet.sheet_view.showGridLines = False

    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = (
        worksheet.PAPERSIZE_A4
    )
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0

    worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    worksheet.print_title_rows = (
        f"1:{header_row}"
    )

    # ---------------------------------------------------------
    # RESPONSE
    # ---------------------------------------------------------

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = (
        f'attachment; filename='
        f'"master_result_{class_level.name}.xlsx"'
    )

    return response

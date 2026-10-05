from django.core.management.base import BaseCommand
from django.db import transaction
from openpyxl import load_workbook

from school.models import Student, ClassLevel, Enrollment, Term


class Command(BaseCommand):
    help = "Import students from an Excel .xlsx file."

    def add_arguments(self, parser):
        parser.add_argument(
            "excel_file",
            type=str,
            help="Path to the Excel file"
        )

    def handle(self, *args, **options):
        excel_file = options["excel_file"]

        term = Term.objects.filter(
            is_current=True
        ).select_related("session").first()

        if not term:
            self.stdout.write(
                self.style.ERROR(
                    "No current term found. Set a term as current in Django Admin."
                )
            )
            return

        try:
            workbook = load_workbook(
                excel_file,
                read_only=True,
                data_only=True
            )
        except FileNotFoundError:
            self.stdout.write(
                self.style.ERROR(
                    f"File not found: {excel_file}"
                )
            )
            return
        except Exception as exc:
            self.stdout.write(
                self.style.ERROR(
                    f"Could not open Excel file: {exc}"
                )
            )
            return

        sheet = workbook.active

        rows = sheet.iter_rows(values_only=True)

        try:
            headers = next(rows)
        except StopIteration:
            self.stdout.write(
                self.style.ERROR("The Excel file is empty.")
            )
            return

        headers = [
            str(h).strip() if h is not None else ""
            for h in headers
        ]

        required = [
            "Admission No.",
            "Surname",
            "First Name",
            "Other Names",
            "Gender",
            "Class",
        ]

        missing = [h for h in required if h not in headers]

        if missing:
            self.stdout.write(
                self.style.ERROR(
                    "Missing required columns: "
                    + ", ".join(missing)
                )
            )
            workbook.close()
            return

        column_index = {
            name: headers.index(name)
            for name in required
        }

        created = 0
        updated = 0
        enrollments_created = 0
        errors = []

        for row_number, row in enumerate(rows, start=2):

            try:
                def value(column):
                    index = column_index[column]

                    if index >= len(row):
                        return ""

                    result = row[index]

                    if result is None:
                        return ""

                    return str(result).strip()

                admission_number = value("Admission No.")
                surname = value("Surname")
                first_name = value("First Name")
                other_names = value("Other Names")
                gender = value("Gender").title()
                class_name = value("Class")

                if not any([
                    admission_number,
                    surname,
                    first_name,
                    other_names,
                    gender,
                    class_name,
                ]):
                    continue

                if not admission_number:
                    raise ValueError("Admission No. is empty.")

                if not surname:
                    raise ValueError("Surname is empty.")

                if not first_name:
                    raise ValueError("First Name is empty.")

                if gender not in ("Male", "Female"):
                    raise ValueError(
                        "Gender must be Male or Female."
                    )

                class_level = ClassLevel.objects.filter(
                    name=class_name
                ).first()

                if not class_level:
                    raise ValueError(
                        f"Class '{class_name}' does not exist."
                    )

                with transaction.atomic():

                    student = Student.objects.filter(
                        admission_number=admission_number
                    ).first()

                    if student:
                        student.surname = surname
                        student.first_name = first_name
                        student.other_names = other_names
                        student.gender = gender
                        student.current_class = class_level
                        student.active = True
                        student.save()

                        updated += 1

                    else:
                        student = Student.objects.create(
                            admission_number=admission_number,
                            surname=surname,
                            first_name=first_name,
                            other_names=other_names,
                            gender=gender,
                            current_class=class_level,
                            active=True,
                        )

                        created += 1

                    enrollment, was_created = Enrollment.objects.get_or_create(
                        student=student,
                        term=term,
                        defaults={
                            "class_level": class_level
                        }
                    )

                    if not was_created:
                        if enrollment.class_level != class_level:
                            enrollment.class_level = class_level
                            enrollment.save()
                    else:
                        enrollments_created += 1

            except Exception as exc:
                errors.append(
                    f"Row {row_number}: {exc}"
                )

        workbook.close()

        self.stdout.write("")
        self.stdout.write(
            self.style.SUCCESS("EXCEL IMPORT COMPLETE")
        )
        self.stdout.write("------------------------------")
        self.stdout.write(f"Students created: {created}")
        self.stdout.write(f"Students updated: {updated}")
        self.stdout.write(
            f"Enrollments created: {enrollments_created}"
        )
        self.stdout.write(f"Errors: {len(errors)}")

        if errors:
            self.stdout.write("")
            self.stdout.write(
                self.style.WARNING("IMPORT ERRORS:")
            )

            for error in errors:
                self.stdout.write(
                    self.style.WARNING(error)
                )
